import io

from core.doctext import extract_text


def tiny_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
            b"/Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for i, o in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + o + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer << /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return out.getvalue()


def test_pdf():
    assert "Week 9 Quiz 2" in extract_text("Teaching Plan.pdf", tiny_pdf("Week 9 Quiz 2"))


def test_docx():
    import docx
    d = docx.Document()
    d.add_paragraph("CA1 presentation in Week 12")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Week 9"
    t.rows[0].cells[1].text = "Quiz 2"
    buf = io.BytesIO()
    d.save(buf)
    text = extract_text("plan.DOCX", buf.getvalue())
    assert "CA1 presentation in Week 12" in text and "Week 9 | Quiz 2" in text


def test_pptx():
    from pptx import Presentation
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Assessment schedule"
    s.placeholders[1].text = "Test on 10 Nov"
    buf = io.BytesIO()
    p.save(buf)
    text = extract_text("w1.pptx", buf.getvalue())
    assert "Assessment schedule" in text and "Test on 10 Nov" in text


def test_txt_html_and_unknown():
    assert extract_text("notes.txt", "Due Friday".encode()) == "Due Friday"
    assert extract_text("page.html", b"<p>Quiz&nbsp;<b>Friday</b></p>") == "Quiz Friday"
    assert extract_text("video.mp4", b"\x00\x01") == ""
    assert extract_text("broken.pdf", b"not a pdf") == ""
