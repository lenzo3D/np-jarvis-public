"""Plain text from course files (PDF, DOCX, PPTX, TXT, HTML). Never raises: unreadable → ""."""
from __future__ import annotations

import io
import logging
from pathlib import Path

from core.text import html_to_text

log = logging.getLogger(__name__)
MAX_CHARS = 400_000


def _pdf(data: bytes) -> str:
    import pdfplumber

    parts = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages[:150]:
            parts.append(page.extract_text() or "")
    return "\n".join(parts)


def _docx(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(parts)


def _pptx(data: bytes) -> str:
    from pptx import Presentation

    parts = []
    for i, slide in enumerate(Presentation(io.BytesIO(data)).slides, 1):
        parts.append(f"--- slide {i}")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def extract_text(name: str, data: bytes) -> str:
    ext = Path(name).suffix.lower()
    try:
        if ext == ".pdf":
            text = _pdf(data)
        elif ext == ".docx":
            text = _docx(data)
        elif ext == ".pptx":
            text = _pptx(data)
        elif ext in (".txt", ".md", ".csv"):
            text = data.decode("utf-8", errors="replace")
        elif ext in (".html", ".htm"):
            text = html_to_text(data.decode("utf-8", errors="replace"), limit=MAX_CHARS)
        else:
            return ""
    except Exception as e:  # noqa: BLE001 - corrupt or unsupported file
        log.info("could not read %s: %s", name, e)
        return ""
    return text.strip()[:MAX_CHARS]
